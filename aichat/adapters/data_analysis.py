"""数据分析适配器 — 把 data-analyst-tool 的 AI 数据分析工作流程接入 AIChat。

data-analyst-tool 是一个 AI 驱动的 Excel/CSV 数据分析系统，核心流程：
  1. 加载数据文件（Excel/CSV/Word）→ 提取表头 + 样本 + 统计概览
  2. LLM 规划分析维度（根据数据说明 + 分析重点）
  3. LLM 生成 pandas 脚本 → 执行 → 产出统计摘要（numeric_stats / categorical_stats）
  4. LLM 深度分析（洞察 / 发现 / 结论）
  5. 生成 HTML 报告（KPI 卡片 + ECharts 图表）

本适配器把这个流程桥接到 AIChat：
  - 数据文件加载通过 <|tool|> 协议调用（load_spreadsheet / run_pandas 等）
  - 分析结果通过 <|memory|> 协议存入四层记忆
  - 多次分析的历史可跨会话检索
  - 不修改原 data-analyst-tool 项目代码
"""
from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

from ..protocol.types import Message, Session, MemoryResult, ChatRole
from ..protocol.adapter import ChatAdapter


class DataAnalysisAdapter(ChatAdapter):
    """数据分析适配器。

    用法：
        adapter = DataAnalysisAdapter(data_file="/path/to/data.xlsx")
        # 或通过 kwargs
        adapter = DataAnalysisAdapter(data_file="...", data_description="2024年销售数据")
    """

    def __init__(self, **kwargs):
        self._data_file: str = kwargs.get("data_file", "")
        self._data_description: str = kwargs.get("data_description", "")
        self._analysis_focus: str = kwargs.get("analysis_focus", "")
        self._file_info: dict = {}
        self._load_file_info()

    # ============================================================
    # ChatAdapter 接口
    # ============================================================

    @property
    def key(self) -> str:
        return "data_analysis"

    @property
    def display_name(self) -> str:
        if self._data_file:
            name = Path(self._data_file).stem
            return f"📊 数据分析 - {name}"
        return "📊 数据分析"

    @property
    def description(self) -> str:
        return "AI 驱动的 Excel/CSV 数据分析：规划→统计→洞察→报告"

    def build_system_prompt(self, session: Session) -> str:
        """构建数据分析系统提示词。"""
        from ..tools import TOOL_PROTOCOL_MD, MEMORY_PROTOCOL_MD

        parts = [
            "你是一个专业的数据分析顾问，服务于 AIChat 数据分析模式。",
            "你可以帮助用户分析 Excel/CSV 数据文件，提供洞察和建议。",
            "",
            "工作原则：",
            "1. 先用 load_spreadsheet 工具加载数据文件，了解表头和数据概况",
            "2. 根据用户的分析重点，用 run_pandas 执行统计计算",
            "3. 拿到统计结果后做深度分析，给出洞察和结论",
            "4. 用 generate_report 可以生成 HTML 报告（含图表）",
            "5. 回答要结构化：核心结论先行，数据支撑在后",
            "6. 数字要精确，引用统计结果时要标注来源",
            "7. 发现异常或趋势时主动指出，不要等用户问",
            "8. 对缺失数据/异常值要提醒用户",
        ]

        # 注入数据文件上下文
        ctx = self._build_data_context()
        if ctx:
            parts.append("")
            parts.append("# 当前数据文件上下文")
            parts.append(ctx)

        # 用户的数据说明和分析重点
        if self._data_description:
            parts.append(f"\n## 数据说明（用户提供）\n{self._data_description}")
        if self._analysis_focus:
            parts.append(f"\n## 分析重点（用户指定）\n{self._analysis_focus}")

        # 工具协议 + 记忆协议 + 流程蓝图协议
        parts.append("")
        parts.append(TOOL_PROTOCOL_MD)
        parts.append("")
        parts.append(MEMORY_PROTOCOL_MD)
        try:
            from ..flow import get_flow_protocol_md
            parts.append("")
            parts.append(get_flow_protocol_md())
        except Exception:
            pass

        return "\n".join(parts)

    def build_context_messages(self, session: Session,
                                memory_results: list[MemoryResult]) -> list[Message]:
        """注入历史分析记忆。"""
        msgs: list[Message] = []
        if memory_results:
            facts: list[str] = []
            for r in memory_results[:10]:
                line = f"- [{r.scope}] {r.content[:200]}"
                facts.append(line)
            if facts:
                msgs.append(Message(
                    role=ChatRole.SYSTEM,
                    content="以下是相关的历史分析记忆：\n\n" + "\n".join(facts)
                ))
        return msgs

    def get_scope(self, session: Session) -> str:
        """返回数据分析的记忆 scope。"""
        if self._data_file:
            name = Path(self._data_file).stem
            return f"data_analysis:{name}"
        return "data_analysis"

    def get_memory_filters(self, session: Session) -> dict[str, Any]:
        """返回记忆过滤条件。"""
        return {"adapter": "data_analysis"}

    def on_response(self, session: Session, response: str) -> dict[str, Any] | None:
        """检测 AI 响应中的报告生成指令。"""
        # 检测是否生成了 HTML 报告
        if "<html" in response.lower() or "<!DOCTYPE html" in response.lower():
            return {"type": "report", "data": {"format": "html", "content": response}}
        # 检测 JSON 格式的分析结论
        if "```json" in response and '"findings"' in response:
            try:
                start = response.index("```json") + 7
                end = response.index("```", start)
                data = json.loads(response[start:end].strip())
                return {"type": "analysis_result", "data": data}
            except (ValueError, json.JSONDecodeError):
                pass
        return None

    # ============================================================
    # 数据文件加载
    # ============================================================

    def _load_file_info(self) -> None:
        """加载数据文件的基本信息。"""
        if not self._data_file or not os.path.exists(self._data_file):
            return

        p = Path(self._data_file)
        info: dict[str, Any] = {
            "file_name": p.name,
            "file_path": str(p),
            "file_size": p.stat().st_size,
            "suffix": p.suffix.lower(),
        }

        # 根据文件类型加载信息
        if p.suffix.lower() in (".xlsx", ".xls"):
            info.update(self._load_excel_info(p))
        elif p.suffix.lower() == ".csv":
            info.update(self._load_csv_info(p))
        elif p.suffix.lower() in (".docx", ".doc"):
            info.update(self._load_docx_info(p))

        self._file_info = info

    def _load_excel_info(self, path: Path) -> dict:
        """加载 Excel 文件信息：sheet 名 + 表头 + 样本行 + 行数。"""
        info: dict[str, Any] = {"type": "excel"}
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            sheets: list[dict] = []
            for sn in wb.sheetnames:
                ws = wb[sn]
                rows = list(ws.iter_rows(max_row=5, values_only=True))
                headers = [str(c) if c is not None else "" for c in rows[0]] if rows else []
                sample = [list(r) for r in rows[1:4]] if len(rows) > 1 else []
                sheets.append({
                    "name": sn,
                    "headers": headers,
                    "sample_rows": sample,
                    "row_count_estimate": ws.max_row,
                    "col_count": len(headers),
                })
            wb.close()
            info["sheets"] = sheets
            info["sheet_names"] = [s["name"] for s in sheets]
        except ImportError:
            info["error"] = "需要 openpyxl 库来读取 Excel 文件"
        except Exception as e:
            info["error"] = str(e)
        return info

    def _load_csv_info(self, path: Path) -> dict:
        """加载 CSV 文件信息。"""
        info: dict[str, Any] = {"type": "csv"}
        try:
            import csv
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                rows = list(reader)
            headers = [str(c) for c in rows[0]] if rows else []
            sample = [list(r) for r in rows[1:4]] if len(rows) > 1 else []
            info["headers"] = headers
            info["sample_rows"] = sample
            info["row_count"] = len(rows) - 1
            info["col_count"] = len(headers)
        except Exception as e:
            info["error"] = str(e)
        return info

    def _load_docx_info(self, path: Path) -> dict:
        """加载 Word 文件信息。"""
        info: dict[str, Any] = {"type": "docx"}
        try:
            from docx import Document
            doc = Document(str(path))
            info["paragraph_count"] = len(doc.paragraphs)
            info["table_count"] = len(doc.tables)
            # 取前5段文本作为摘要
            info["preview"] = [p.text[:100] for p in doc.paragraphs[:5]]
        except ImportError:
            info["error"] = "需要 python-docx 库来读取 Word 文件"
        except Exception as e:
            info["error"] = str(e)
        return info

    # ============================================================
    # 上下文组装
    # ============================================================

    def _build_data_context(self) -> str:
        """构建注入 system prompt 的数据上下文。"""
        if not self._file_info:
            return ""

        info = self._file_info
        parts: list[str] = []

        parts.append(f"文件名：{info.get('file_name', '?')}")
        parts.append(f"文件大小：{info.get('file_size', 0)} bytes")
        parts.append(f"文件类型：{info.get('type', info.get('suffix', '?'))}")

        # Excel 多 sheet
        sheets = info.get("sheets", [])
        if sheets:
            for s in sheets[:3]:  # 最多展示3个sheet
                parts.append(f"\n## Sheet: {s['name']}")
                parts.append(f"行数: {s.get('row_count_estimate', '?')} | 列数: {s.get('col_count', '?')}")
                headers = s.get("headers", [])
                if headers:
                    parts.append(f"表头: {', '.join(str(h) for h in headers[:20])}")
                sample = s.get("sample_rows", [])
                if sample:
                    parts.append("样本行（前3行）:")
                    for i, row in enumerate(sample[:3], 1):
                        vals = [str(v)[:30] for v in row[:10]]
                        parts.append(f"  行{i}: {', '.join(vals)}")

        # CSV 单表
        elif info.get("headers"):
            parts.append(f"\n## 数据表")
            parts.append(f"行数: {info.get('row_count', '?')} | 列数: {info.get('col_count', '?')}")
            headers = info.get("headers", [])
            parts.append(f"表头: {', '.join(str(h) for h in headers[:20])}")
            sample = info.get("sample_rows", [])
            if sample:
                parts.append("样本行:")
                for i, row in enumerate(sample[:3], 1):
                    vals = [str(v)[:30] for v in row[:10]]
                    parts.append(f"  行{i}: {', '.join(vals)}")

        # Word 文档
        elif info.get("type") == "docx":
            parts.append(f"段落数: {info.get('paragraph_count', '?')}")
            parts.append(f"表格数: {info.get('table_count', '?')}")
            preview = info.get("preview", [])
            if preview:
                parts.append("文档预览:")
                for p in preview:
                    parts.append(f"  {p}")

        # 错误信息
        if info.get("error"):
            parts.append(f"\n⚠️ 文件加载错误: {info['error']}")

        # 限制总量
        result = "\n".join(parts)
        if len(result) > 4000:
            result = result[:4000] + "\n... [上下文截断，完整数据可通过工具读取]"
        return result

    # ============================================================
    # 工具支持：提供给 skill 调用的辅助方法
    # ============================================================

    def get_file_path(self) -> str:
        return self._data_file

    def get_file_info(self) -> dict:
        return self._file_info

    def get_sheet_names(self) -> list[str]:
        return self._file_info.get("sheet_names", [])

    def get_headers(self, sheet_name: str | None = None) -> list[str]:
        sheets = self._file_info.get("sheets", [])
        if sheet_name:
            for s in sheets:
                if s["name"] == sheet_name:
                    return s.get("headers", [])
        if sheets:
            return sheets[0].get("headers", [])
        return self._file_info.get("headers", [])
